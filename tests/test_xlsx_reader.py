"""XLSX BoQ reader (docs/ui-spec.md §4, U1): same lines as the CSV reader, values only."""

import csv
import io
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from oris_matcher.io import xlsx_reader
from oris_matcher.io.boq_reader import BoqFormatError, MissingColumnsError, read_boq
from oris_matcher.io.xlsx_reader import (
    EmptyBoqError,
    TooManyRowsError,
    UnsupportedTypeError,
    WorkbookTooLargeError,
    read_upload,
    read_xlsx,
    render_cell,
)

ROOT = Path(__file__).resolve().parents[1]
EN_INPUT = ROOT / "input" / "boq_dataset_input_en.csv"
FR_INPUT = ROOT / "input" / "boq_dataset_input_fr.csv"
HEADER = ["Item No.", "Short Description", "Long Description", "Unit", "BoQ Qty"]
CSV_HEADER = ",".join(HEADER) + "\r\n"


def csv_records(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def workbook_bytes(*sheets: tuple[str, list[list[object]], str]) -> bytes:
    """Build an .xlsx with the given (title, rows, state) sheets, in order."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows, state in sheets:
        sheet = book.create_sheet(title)
        sheet.sheet_state = state
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def xlsx_of(path: Path) -> bytes:
    return workbook_bytes(("BoQ", [list(r) for r in csv_records(path)], "visible"))


@pytest.mark.parametrize("path", [EN_INPUT, FR_INPUT])
def test_u1_xlsx_gives_the_csv_line_ids_and_cells(path: Path) -> None:
    from_csv = read_boq(path)
    from_xlsx = read_xlsx(xlsx_of(path))

    assert [line.line_id for line in from_xlsx.lines] == [line.line_id for line in from_csv.lines]
    assert [line.raw_row for line in from_xlsx.lines] == [line.raw_row for line in from_csv.lines]
    assert [line.kind for line in from_xlsx.lines] == [line.kind for line in from_csv.lines]
    assert [line.section_path for line in from_xlsx.lines] == [
        line.section_path for line in from_csv.lines
    ]
    assert from_xlsx.header == from_csv.header


def test_numeric_cells_keep_their_number_format() -> None:
    rows: list[list[object]] = [
        HEADER,
        [9, "Concrete", "", "m3", 12],
        [9.01, "Steel", "", "t", 2.5],
    ]
    data = workbook_bytes(("BoQ", rows, "visible"))
    book = openpyxl.load_workbook(io.BytesIO(data))
    book["BoQ"]["A2"].number_format = "00"
    book["BoQ"]["A3"].number_format = "00.00"
    book["BoQ"]["E3"].number_format = "0.000"
    buffer = io.BytesIO()
    book.save(buffer)

    lines = read_xlsx(buffer.getvalue()).lines

    assert [line.item_no for line in lines] == ["09", "09.01"]
    assert [line.qty for line in lines] == ["12", "2.500"]


@pytest.mark.parametrize(
    ("value", "number_format", "expected"),
    [
        (12, "General", "12"),
        (12.0, "General", "12"),
        (2.5, "General", "2.5"),
        (1234.5, "#,##0.00", "1,234.50"),
        (7, "000", "007"),
        (True, "General", "TRUE"),
        (None, "General", ""),
        ("09.01.0010.", "@", "09.01.0010."),
        (datetime(2026, 10, 9), "yyyy-mm-dd", "2026-10-09"),
    ],
)
def test_render_cell(value: object, number_format: str, expected: str) -> None:
    assert render_cell(value, number_format) == expected


def test_the_first_visible_sheet_is_read() -> None:
    hidden = [["not", "a", "boq"]]
    rows: list[list[object]] = [HEADER, ["01.01.0010.", "Concrete", "", "m3", "4"]]
    data = workbook_bytes(("Hidden", hidden, "hidden"), ("BoQ", rows, "visible"))

    assert [line.item_no for line in read_xlsx(data).lines] == ["01.01.0010."]


def test_trailing_blank_rows_are_dropped_like_trailing_csv_line_breaks() -> None:
    rows: list[list[object]] = [HEADER, ["01.01.0010.", "Concrete", "", "m3", "4"], [], [None]]
    data = workbook_bytes(("BoQ", rows, "visible"))

    assert len(read_xlsx(data).lines) == 1


def test_missing_column_is_reported_by_name() -> None:
    rows: list[list[object]] = [HEADER[:4], ["01", "Concrete", "", "m3"]]

    with pytest.raises(MissingColumnsError) as caught:
        read_xlsx(workbook_bytes(("BoQ", rows, "visible")))

    assert caught.value.missing == ("BoQ Qty",)


def test_too_many_rows_stops_reading() -> None:
    rows: list[list[object]] = [HEADER, *([["01", "x", "", "m", "1"]] * 6)]

    with pytest.raises(TooManyRowsError):
        read_xlsx(workbook_bytes(("BoQ", rows, "visible")), max_rows=5)


def test_a_non_workbook_is_a_format_error() -> None:
    with pytest.raises(BoqFormatError):
        read_xlsx(b"PK\x03\x04 not really a zip")


@pytest.mark.parametrize("name", ["boq.xlsm", "boq.xls", "boq.txt", "boq"])
def test_unsupported_types_are_refused(name: str) -> None:
    with pytest.raises(BoqFormatError, match="unsupported"):
        read_upload(name, b"Item No.,x\n")


def test_read_upload_dispatches_on_the_extension() -> None:
    data = EN_INPUT.read_bytes()

    assert read_upload("BOQ.CSV", data).boq == read_boq(data)
    assert read_upload("boq.xlsx", xlsx_of(EN_INPUT)).boq.lines[0].line_id == (
        read_boq(data).lines[0].line_id
    )


def test_a_cp1252_csv_is_read_with_its_warning() -> None:
    data = CSV_HEADER + "1,Béton,,m3,4\r\n"

    uploaded = read_upload("boq.csv", data.encode("cp1252"))

    assert uploaded.boq.encoding == "cp1252"
    assert len(uploaded.warnings) == 1
    assert "cp1252" in uploaded.warnings[0]


@pytest.mark.parametrize("data", [b"", b" \r\n", CSV_HEADER.encode()])
def test_an_empty_csv_is_refused_as_empty(data: bytes) -> None:
    with pytest.raises(EmptyBoqError):
        read_upload("boq.csv", data)


def test_a_csv_over_the_row_limit_is_refused() -> None:
    with pytest.raises(TooManyRowsError):
        read_upload("boq.csv", EN_INPUT.read_bytes(), max_rows=10)


def zip_with(member: str, data: bytes) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(member, data)
    return buffer.getvalue()


def test_a_member_that_inflates_far_beyond_its_size_is_refused_before_parsing() -> None:
    data = zip_with("xl/sharedStrings.xml", b"a" * (8 * 1024 * 1024))

    with pytest.raises(WorkbookTooLargeError):
        read_xlsx(data)


def test_a_workbook_over_the_uncompressed_cap_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    data = xlsx_of(EN_INPUT)
    monkeypatch.setattr(xlsx_reader, "MAX_XLSX_UNCOMPRESSED_BYTES", 1024)

    with pytest.raises(WorkbookTooLargeError):
        read_xlsx(data)


def test_an_ordinary_workbook_passes_the_inflation_guard() -> None:
    assert read_xlsx(xlsx_of(FR_INPUT)).lines


def test_xlsx_is_refused_when_openpyxl_parses_without_defusedxml(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import openpyxl.xml  # noqa: PLC0415

    monkeypatch.setattr(openpyxl.xml, "DEFUSEDXML", False)

    with pytest.raises(UnsupportedTypeError, match="defusedxml"):
        read_xlsx(xlsx_of(EN_INPUT))
